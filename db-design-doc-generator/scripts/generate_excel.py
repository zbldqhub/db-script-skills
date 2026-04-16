import psycopg2
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import re
import os
import json
import argparse

HEADERS = [
    'tabname', 'namec', 'colname', 'datatype', 'disporder', 'disptype',
    'nullable', 'newedit', 'editable', 'indexname', 'indextype', 'indexcols',
    'major_minor', 'reference', '字典值', '字段说明'
]

def format_datatype(data_type, char_len, num_prec, num_scale, column_default):
    if column_default and 'nextval' in str(column_default):
        if data_type == 'integer':
            return 'serial'
        elif data_type == 'smallint':
            return 'smallserial'
        elif data_type == 'bigint':
            return 'bigserial'
    if data_type == 'character varying':
        return f'varchar({char_len})' if char_len else 'varchar'
    elif data_type == 'character':
        return f'char({char_len})' if char_len else 'char'
    elif data_type == 'numeric':
        if num_prec is not None and num_scale is not None:
            return f'numeric({num_prec},{num_scale})'
        elif num_prec is not None:
            return f'numeric({num_prec})'
        else:
            return 'numeric'
    elif data_type == 'timestamp without time zone':
        return 'timestamp'
    elif data_type == 'timestamp with time zone':
        return 'timestamptz'
    elif data_type == 'double precision':
        return 'double precision'
    else:
        return data_type

def parse_index(indexdef):
    m = re.search(r'CREATE\s+(UNIQUE\s+)?INDEX\s+(\w+)\s+ON\s+\w+\.\w+\s+USING\s+\w+\s+\((.+)\)', indexdef)
    if m:
        is_unique = m.group(1) is not None
        indexname = m.group(2)
        indexcols = m.group(3)
        indextype = 'UNIQUE INDEX' if is_unique else 'INDEX'
        return indexname, indextype, indexcols
    return None, None, None

def get_memo(memo, description, ms):
    if memo and str(memo).strip():
        return str(memo).strip()
    if description and str(description).strip():
        return str(description).strip()
    if ms and str(ms).strip():
        return str(ms).strip()
    return ''

def main():
    parser = argparse.ArgumentParser(description='Generate database design Excel documents from PostgreSQL.')
    parser.add_argument('--db-url', required=True, help='PostgreSQL connection URL')
    parser.add_argument('--output-dir', required=True, help='Output directory for Excel files')
    parser.add_argument('--schema', default='zgis', help='Database schema (default: zgis)')
    parser.add_argument('--major-map', help='Optional JSON file mapping major numbers to filenames')
    args = parser.parse_args()

    db_url = args.db_url
    output_dir = args.output_dir
    schema = args.schema

    major_to_file = {}
    if args.major_map:
        with open(args.major_map, 'r', encoding='utf-8') as f:
            major_to_file = {int(k): v for k, v in json.load(f).items()}

    print("Connecting to database...")
    conn = psycopg2.connect(db_url)
    cursor = conn.cursor()

    # 1. entities
    cursor.execute("SELECT name, namec, major, minor FROM cx_entity WHERE major > 0 ORDER BY major, minor")
    entities = cursor.fetchall()
    entity_map = {}
    tables_by_major = {}
    for name, namec, major, minor in entities:
        entity_map[name] = (namec, major, minor)
        tables_by_major.setdefault(major, []).append(name)

    # Build major_to_file fallback
    for major in tables_by_major.keys():
        if major not in major_to_file:
            major_to_file[major] = str(major)

    # 2. actual columns
    cursor.execute("""
        SELECT table_name, column_name, data_type, character_maximum_length,
               numeric_precision, numeric_scale, is_nullable, ordinal_position, column_default
        FROM information_schema.columns
        WHERE table_schema = %s AND table_name IN (SELECT name FROM cx_entity WHERE major > 0)
        ORDER BY table_name, ordinal_position
    """, (schema,))
    col_map = {}
    for row in cursor.fetchall():
        key = (row[0], row[1])
        col_map[key] = row[2:]

    # 3. cx_fld
    cursor.execute("SELECT tabname, colname, namec, disptype, disporder, newedit, editable, nullable, memo, description, ms FROM cx_fld")
    fld_map = {}
    for row in cursor.fetchall():
        key = (row[0].lower(), row[1].lower())
        if key not in fld_map:
            fld_map[key] = row[2:]

    # 4. indexes
    cursor.execute("SELECT tablename, indexname, indexdef FROM pg_indexes WHERE schemaname = %s", (schema,))
    index_map = {}
    for tablename, indexname, indexdef in cursor.fetchall():
        idx_name, idx_type, idx_cols = parse_index(indexdef)
        if idx_name:
            index_map.setdefault(tablename, []).append((idx_name, idx_type, idx_cols))

    # 5. foreign keys
    cursor.execute("""
        SELECT tc.table_name, kcu.column_name, ccu.table_name, ccu.column_name
        FROM information_schema.table_constraints tc
        JOIN information_schema.key_column_usage kcu ON tc.constraint_name = kcu.constraint_name AND tc.table_schema = kcu.table_schema
        JOIN information_schema.constraint_column_usage ccu ON ccu.constraint_name = tc.constraint_name AND ccu.table_schema = tc.table_schema
        WHERE tc.constraint_type = 'FOREIGN KEY' AND tc.table_schema = %s
    """, (schema,))
    fk_map = {}
    for table_name, column_name, ft, fc in cursor.fetchall():
        fk_map[(table_name, column_name)] = f"{ft}({fc})"

    # 6. dictionary values
    cursor.execute("SELECT tabname, colname, dbvalue, dispc FROM cx_fldvalue ORDER BY tabname, colname, disporder")
    dict_map = {}
    for tabname, colname, dbvalue, dispc in cursor.fetchall():
        dict_map.setdefault((tabname, colname), []).append((dbvalue, dispc))

    cursor.close()
    conn.close()

    col_index_map = {}
    for tablename, idx_list in index_map.items():
        assigned_cols = set()
        for idx_name, idx_type, idx_cols in idx_list:
            cols = [c.strip() for c in idx_cols.split(',')]
            for c in cols:
                if c not in assigned_cols:
                    col_index_map[(tablename, c)] = (idx_name, idx_type, idx_cols)
                    assigned_cols.add(c)
                    break

    header_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
    header_font = Font(bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    os.makedirs(output_dir, exist_ok=True)

    for major, filename_base in sorted(major_to_file.items()):
        tables = tables_by_major.get(major, [])
        if not tables:
            print(f"Skipping {filename_base} (no tables)")
            continue

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for tabname in tables:
            namec, maj, minor = entity_map[tabname]
            major_minor = f"{maj}@{minor}"

            cols = [(tabname, c) for c in [k[1] for k in col_map.keys() if k[0] == tabname]]
            cols.sort(key=lambda k: col_map[k][5])

            base_sheet_name = (namec or tabname)[:31]
            sheet_name = base_sheet_name
            suffix = 1
            while sheet_name in wb.sheetnames:
                suffix_str = f"({suffix})"
                sheet_name = base_sheet_name[:31 - len(suffix_str)] + suffix_str
                suffix += 1
            ws = wb.create_sheet(title=sheet_name)

            for col_idx, h in enumerate(HEADERS, 1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')

            for row_idx, key in enumerate(cols, 2):
                tab, colname = key
                data_type, char_len, num_prec, num_scale, is_nullable, ord_pos, column_default = col_map[key]
                datatype = format_datatype(data_type, char_len, num_prec, num_scale, column_default)

                lookup_key = (tab.lower(), colname.lower())
                if lookup_key in fld_map:
                    f_namec, f_disptype, f_disporder, f_newedit, f_editable, f_nullable, f_memo, f_desc, f_ms = fld_map[lookup_key]
                else:
                    f_namec, f_disptype, f_disporder, f_newedit, f_editable, f_nullable = '', '', '', '', '', ''
                    f_memo, f_desc, f_ms = '', '', ''

                if colname.lower() == 'id':
                    if not f_namec:
                        f_namec = 'ID'
                    f_disporder = 0
                    f_disptype = 1
                    f_nullable = 0
                    f_newedit = 0
                    f_editable = 0

                if f_nullable == '' or f_nullable is None:
                    nullable_val = 0 if is_nullable == 'NO' else 1
                else:
                    nullable_val = f_nullable

                idx_info = col_index_map.get(key)
                if idx_info:
                    idx_name, idx_type, idx_cols = idx_info
                else:
                    idx_name, idx_type, idx_cols = '', '', ''

                reference = fk_map.get(key, '')
                dict_vals = dict_map.get(key, [])
                dict_str = '; '.join([f"{db}={disp}" for db, disp in dict_vals])
                memo_str = get_memo(f_memo, f_desc, f_ms)

                values = [
                    tab, f_namec or '', colname, datatype,
                    f_disporder if f_disporder != '' else '',
                    f_disptype if f_disptype != '' else '',
                    nullable_val,
                    f_newedit if f_newedit != '' else '',
                    f_editable if f_editable != '' else '',
                    idx_name, idx_type, idx_cols, major_minor,
                    reference, dict_str, memo_str,
                ]

                for col_idx, val in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.border = thin_border

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[col_letter].width = adjusted_width

        filepath = os.path.join(output_dir, f"{filename_base}.xlsx")
        wb.save(filepath)
        print(f"Saved: {filepath} ({len(tables)} sheets)")

    print("Done.")

if __name__ == '__main__':
    main()
